"""
Reescribe todos los ingestores con sintaxis correcta.
Ejecutar desde rag_software_intelligence:
  python fix_all_ingestores.py
"""
import os

FILES = {}

# ══════════════════════════════════════════════════════
FILES["ingestors/base_loader.py"] = """\
from abc import ABC, abstractmethod
from typing import List
from langchain.schema import Document


class BaseLoader(ABC):
    \"\"\"Interfaz comun para todos los ingestores del sistema RAG.\"\"\"

    @abstractmethod
    def load(self) -> List[Document]:
        pass

    def _make_doc(self, content: str, **meta) -> Document:
        return Document(page_content=content.strip(), metadata=meta)

    def _safe_read(self, path: str, encoding: str = "utf-8") -> str:
        try:
            with open(path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            with open(path, "r", encoding="latin-1") as f:
                return f.read()
"""

# ══════════════════════════════════════════════════════
FILES["ingestors/git_loader.py"] = """\
import os
from typing import List
import git
from langchain.schema import Document
from ingestors.base_loader import BaseLoader

CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".java", ".cs", ".go", ".rb",
    ".php", ".cpp", ".c", ".h", ".rs", ".kt",
    ".sql", ".sh", ".yaml", ".yml", ".json", ".toml",
    ".md", ".txt", ".cfg", ".ini",
}
IGNORE_DIRS = {".git", "__pycache__", "node_modules",
               ".venv", "venv", "dist", "build"}


class GitLoader(BaseLoader):
    \"\"\"Clona un repositorio Git e indexa su codigo fuente.\"\"\"

    def __init__(self, repo_url: str, clone_dir: str):
        self.repo_url  = repo_url
        self.clone_dir = clone_dir

    def load(self) -> List[Document]:
        repo_path = self._get_repo()
        documents = []

        for root, dirs, files in os.walk(repo_path):
            dirs[:] = [d for d in dirs if d not in IGNORE_DIRS]
            for filename in files:
                ext = os.path.splitext(filename)[1].lower()
                if ext not in CODE_EXTENSIONS:
                    continue
                filepath = os.path.join(root, filename)
                rel_path = os.path.relpath(filepath, repo_path)
                content  = self._safe_read(filepath)
                if not content.strip():
                    continue
                lines = content.splitlines()
                numbered_lines = []
                for i, line in enumerate(lines):
                    numbered_lines.append(str(i + 1).rjust(4) + ": " + line)
                numbered = "\\n".join(numbered_lines)
                doc = self._make_doc(
                    content="# Archivo: " + rel_path + "\\n\\n" + numbered,
                    source=rel_path,
                    file_type=ext,
                    repo_url=self.repo_url,
                    lines_total=len(lines),
                    loader="GitLoader",
                )
                documents.append(doc)

        print("  [GitLoader]", len(documents), "archivos indexados")
        return documents

    def _get_repo(self) -> str:
        repo_name  = self.repo_url.rstrip("/").split("/")[-1].replace(".git", "")
        local_path = os.path.join(self.clone_dir, repo_name)
        if os.path.exists(local_path):
            print("  [GitLoader] Repo encontrado en", local_path)
            try:
                repo = git.Repo(local_path)
                repo.remotes.origin.pull()
            except Exception:
                pass
        else:
            print("  [GitLoader] Clonando", self.repo_url)
            os.makedirs(self.clone_dir, exist_ok=True)
            git.Repo.clone_from(self.repo_url, local_path)
        return local_path
"""

# ══════════════════════════════════════════════════════
FILES["ingestors/drawio_loader.py"] = """\
import os
import re
from typing import List
from xml.etree import ElementTree as ET
from langchain.schema import Document
from ingestors.base_loader import BaseLoader


class DrawioLoader(BaseLoader):
    \"\"\"Parsea diagramas Draw.io y extrae nodos, capas y conexiones.\"\"\"

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        if not os.path.exists(self.file_path):
            raise FileNotFoundError("Archivo no encontrado: " + self.file_path)

        tree     = ET.parse(self.file_path)
        root     = tree.getroot()
        cells    = root.findall(".//mxCell")
        filename = os.path.basename(self.file_path)
        documents = []
        nodes  = {}
        edges  = []
        layers = {}

        for cell in cells:
            cell_id   = cell.get("id", "")
            style     = cell.get("style", "")
            value     = self._clean_html(cell.get("value", ""))
            source_id = cell.get("source", "")
            target_id = cell.get("target", "")
            vertex    = cell.get("vertex", "0") == "1"
            edge      = cell.get("edge",   "0") == "1"

            if "swimlane" in style and value:
                layers[cell_id] = value
            elif vertex and value and cell_id not in ("0", "1"):
                nodes[cell_id] = value
            elif edge and (source_id or target_id):
                edges.append((source_id, target_id, value))

        for layer_id, layer_name in layers.items():
            text = "Capa de arquitectura: " + layer_name
            documents.append(self._make_doc(
                content=text,
                source=filename, file_type=".drawio",
                element_type="layer", element_id=layer_id,
                loader="DrawioLoader",
            ))

        for node_id, label in nodes.items():
            text = "Componente de arquitectura: " + label
            documents.append(self._make_doc(
                content=text,
                source=filename, file_type=".drawio",
                element_type="node", element_id=node_id,
                loader="DrawioLoader",
            ))

        for src_id, tgt_id, label in edges:
            src_label  = nodes.get(src_id) or layers.get(src_id) or src_id
            tgt_label  = nodes.get(tgt_id) or layers.get(tgt_id) or tgt_id
            edge_label = (" [" + label + "]") if label else ""
            text = ("Relacion en arquitectura: "
                    + src_label + " -> " + tgt_label + edge_label)
            documents.append(self._make_doc(
                content=text,
                source=filename, file_type=".drawio",
                element_type="edge",
                element_id=src_id + "->" + tgt_id,
                loader="DrawioLoader",
            ))

        print("  [DrawioLoader]", len(layers), "capas,",
              len(nodes), "nodos,", len(edges), "conexiones ->",
              len(documents), "documentos")
        return documents

    @staticmethod
    def _clean_html(text: str) -> str:
        text = re.sub(r"<br\\s*/?>", " ", text, flags=re.IGNORECASE)
        text = re.sub(r"<[^>]+>", "", text)
        return text.strip()
"""

# ══════════════════════════════════════════════════════
FILES["ingestors/excel_loader.py"] = """\
import os
from typing import List
import openpyxl
from langchain.schema import Document
from ingestors.base_loader import BaseLoader


class ExcelLoader(BaseLoader):
    \"\"\"Lee diccionarios de datos en Excel e indexa tablas y campos.\"\"\"

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        if not os.path.exists(self.file_path):
            raise FileNotFoundError("Archivo no encontrado: " + self.file_path)

        wb       = openpyxl.load_workbook(self.file_path, data_only=True)
        filename = os.path.basename(self.file_path)
        documents = []

        for sheet_name in wb.sheetnames:
            ws   = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue

            header_row_idx = None
            for i, row in enumerate(rows):
                if len([c for c in row if c is not None]) >= 2:
                    header_row_idx = i
                    break
            if header_row_idx is None:
                continue

            headers   = [str(c).strip() if c else "" for c in rows[header_row_idx]]
            data_rows = rows[header_row_idx + 1:]

            resumen = self._make_doc(
                content=("Tabla de base de datos: " + sheet_name
                         + "\\nColumnas: " + ", ".join(h for h in headers if h)
                         + "\\nTotal campos: " + str(len(data_rows))),
                source=filename, file_type=".xlsx",
                sheet=sheet_name, element_type="table_summary",
                loader="ExcelLoader",
            )
            documents.append(resumen)

            for row_idx, row in enumerate(data_rows, start=header_row_idx + 2):
                if all(v is None for v in row):
                    continue
                campo_dict = {}
                for h, val in zip(headers, row):
                    if h and val is not None:
                        campo_dict[h] = str(val).strip()
                if not campo_dict:
                    continue

                field_name = (campo_dict.get("Campo")
                              or campo_dict.get("Field")
                              or list(campo_dict.values())[0])
                lines = ["Campo '" + field_name + "' en tabla '" + sheet_name + "':"]
                for k, v in campo_dict.items():
                    if v and v.lower() not in ("none", "null", "-", "—"):
                        lines.append("  " + k + ": " + v)

                doc = self._make_doc(
                    content="\\n".join(lines),
                    source=filename, file_type=".xlsx",
                    sheet=sheet_name, row=row_idx,
                    field_name=field_name, element_type="field",
                    loader="ExcelLoader",
                )
                documents.append(doc)

                fk_val = campo_dict.get("FK") or campo_dict.get("Relaciones") or ""
                if fk_val and fk_val not in ("—", "-", ""):
                    rel_doc = self._make_doc(
                        content=("FK: tabla '" + sheet_name
                                 + "', campo '" + field_name
                                 + "' referencia a " + fk_val),
                        source=filename, file_type=".xlsx",
                        sheet=sheet_name, field_name=field_name,
                        element_type="foreign_key", loader="ExcelLoader",
                    )
                    documents.append(rel_doc)

        print("  [ExcelLoader]", len(documents), "documentos desde", filename)
        return documents
"""

# ══════════════════════════════════════════════════════
FILES["ingestors/doc_loader.py"] = """\
import os
import re
from typing import List
from langchain.schema import Document
from ingestors.base_loader import BaseLoader


class DocLoader(BaseLoader):
    \"\"\"Carga documentacion tecnica en PDF o Markdown.\"\"\"

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        if not os.path.exists(self.file_path):
            raise FileNotFoundError("Documento no encontrado: " + self.file_path)
        ext = os.path.splitext(self.file_path)[1].lower()
        if ext == ".pdf":
            return self._load_pdf()
        elif ext in (".md", ".markdown", ".txt"):
            return self._load_markdown()
        else:
            raise ValueError("Extension no soportada: " + ext)

    def _load_pdf(self) -> List[Document]:
        from pypdf import PdfReader
        reader   = PdfReader(self.file_path)
        filename = os.path.basename(self.file_path)
        documents = []
        for page_num, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            header = "[PDF: " + filename + " - Pagina " + str(page_num) + "]\\n\\n"
            documents.append(self._make_doc(
                content=header + text,
                source=filename, file_type=".pdf",
                page=page_num, loader="DocLoader",
            ))
        print("  [DocLoader/PDF]", len(documents), "paginas desde", filename)
        return documents

    def _load_markdown(self) -> List[Document]:
        filename  = os.path.basename(self.file_path)
        content   = self._safe_read(self.file_path)
        documents = []
        sections  = re.split(r"(?m)^(#{1,2}\\s+.+)$", content)
        current_title = filename
        buffer = []
        for part in sections:
            if re.match(r"^#{1,2}\\s+", part):
                if buffer:
                    text = "\\n".join(buffer).strip()
                    if text:
                        header = "[MD: " + filename + " - " + current_title + "]\\n\\n"
                        documents.append(self._make_doc(
                            content=header + text,
                            source=filename, file_type=".md",
                            section=current_title, loader="DocLoader",
                        ))
                current_title = part.strip("# ").strip()
                buffer = []
            else:
                buffer.append(part)
        if buffer:
            text = "\\n".join(buffer).strip()
            if text:
                header = "[MD: " + filename + " - " + current_title + "]\\n\\n"
                documents.append(self._make_doc(
                    content=header + text,
                    source=filename, file_type=".md",
                    section=current_title, loader="DocLoader",
                ))
        print("  [DocLoader/MD]", len(documents), "secciones desde", filename)
        return documents
"""

# ══════════════════════════════════════════════════════
FILES["ingestors/infra_loader.py"] = """\
import os
from typing import List
from langchain.schema import Document
from ingestors.base_loader import BaseLoader


class InfraLoader(BaseLoader):
    \"\"\"Lee archivos de infraestructura: Dockerfile, docker-compose, K8s.\"\"\"

    def __init__(self, path: str):
        self.path = path

    def load(self) -> List[Document]:
        documents = []
        if os.path.isfile(self.path):
            doc = self._load_file(self.path)
            if doc:
                documents.append(doc)
        elif os.path.isdir(self.path):
            for root, _, files in os.walk(self.path):
                for fname in files:
                    if self._is_infra_file(fname):
                        doc = self._load_file(os.path.join(root, fname))
                        if doc:
                            documents.append(doc)
        print("  [InfraLoader]", len(documents), "archivos de infraestructura")
        return documents

    def _load_file(self, filepath: str):
        filename = os.path.basename(filepath)
        content  = self._safe_read(filepath)
        if not content.strip():
            return None
        infra_type = self._detect_type(filename, content)
        header = "[Infraestructura - " + infra_type + "]\\nArchivo: " + filename + "\\n\\n"
        return self._make_doc(
            content=header + content,
            source=filename,
            file_type=os.path.splitext(filename)[1] or filename,
            infra_type=infra_type, loader="InfraLoader",
        )

    @staticmethod
    def _is_infra_file(filename: str) -> bool:
        name = filename.lower()
        if name in ("dockerfile", ".env.example", "makefile"):
            return True
        ext = os.path.splitext(filename)[1].lower()
        return ext in (".yaml", ".yml", ".toml", ".env")

    @staticmethod
    def _detect_type(filename: str, content: str) -> str:
        name = filename.lower()
        if "dockerfile" in name:
            return "Dockerfile"
        if "docker-compose" in name:
            return "Docker Compose"
        if name.endswith((".yaml", ".yml")):
            return "YAML Config"
        if name.endswith(".toml"):
            return "TOML Config"
        if ".env" in name:
            return "Variables de Entorno"
        return "Archivo de Infraestructura"
"""

# ══════════════════════════════════════════════════════
for ruta, contenido in FILES.items():
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(contenido)
    print("  Corregido:", ruta)

print("\nTodos los ingestores corregidos.")
print("Ejecuta: python tests/test_ingestores_manual.py")
