import os
from typing import List
import openpyxl
from langchain.schema import Document
from ingestors.base_loader import BaseLoader


class ExcelLoader(BaseLoader):
    """Lee diccionarios de datos en Excel e indexa tablas y campos."""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        if not os.path.exists(self.file_path):
            raise FileNotFoundError("Archivo no encontrado: " + self.file_path)

        wb       = openpyxl.load_workbook(self.file_path, data_only=True)
        filename = os.path.basename(self.file_path)
        documents = []

        for sheet_name in wb.sheetnames:
            ws   = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue

            header_row_idx = None
            for i, row in enumerate(rows):
                if len([c for c in row if c is not None]) >= 2:
                    header_row_idx = i
                    break
            if header_row_idx is None:
                continue

            headers   = [str(c).strip() if c else "" for c in rows[header_row_idx]]
            data_rows = rows[header_row_idx + 1:]

            resumen = self._make_doc(
                content=("Tabla de base de datos: " + sheet_name
                         + "\nColumnas: " + ", ".join(h for h in headers if h)
                         + "\nTotal campos: " + str(len(data_rows))),
                source=filename, file_type=".xlsx",
                sheet=sheet_name, element_type="table_summary",
                loader="ExcelLoader",
            )
            documents.append(resumen)

            for row_idx, row in enumerate(data_rows, start=header_row_idx + 2):
                if all(v is None for v in row):
                    continue
                campo_dict = {}
                for h, val in zip(headers, row):
                    if h and val is not None:
                        campo_dict[h] = str(val).strip()
                if not campo_dict:
                    continue

                field_name = (campo_dict.get("Campo")
                              or campo_dict.get("Field")
                              or list(campo_dict.values())[0])
                lines = ["Campo '" + field_name + "' en tabla '" + sheet_name + "':"]
                for k, v in campo_dict.items():
                    if v and v.lower() not in ("none", "null", "-", "—"):
                        lines.append("  " + k + ": " + v)

                doc = self._make_doc(
                    content="\n".join(lines),
                    source=filename, file_type=".xlsx",
                    sheet=sheet_name, row=row_idx,
                    field_name=field_name, element_type="field",
                    loader="ExcelLoader",
                )
                documents.append(doc)

                fk_val = campo_dict.get("FK") or campo_dict.get("Relaciones") or ""
                if fk_val and fk_val not in ("—", "-", ""):
                    rel_doc = self._make_doc(
                        content=("FK: tabla '" + sheet_name
                                 + "', campo '" + field_name
                                 + "' referencia a " + fk_val),
                        source=filename, file_type=".xlsx",
                        sheet=sheet_name, field_name=field_name,
                        element_type="foreign_key", loader="ExcelLoader",
                    )
                    documents.append(rel_doc)

        print("  [ExcelLoader]", len(documents), "documentos desde", filename)
        return documents
